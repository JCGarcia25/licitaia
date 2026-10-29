"""Exportador del Agente 4: checklist y matrices a DOCX / XLSX / PDF."""
from __future__ import annotations

import io
from datetime import datetime


def _formato_cop(valor: float | None) -> str:
    if valor is None:
        return "N/D"
    return f"$ {valor:,.0f} COP"


def export_docx(titulo: str, secciones: list[tuple[str, list[list[str]]]]) -> bytes:
    """Genera un .docx con secciones y tablas (carta, checklist, matriz)."""
    import docx

    doc = docx.Document()
    doc.add_heading(titulo, level=0)
    doc.add_paragraph(f"Generado por LicitaIA · {datetime.utcnow():%d/%m/%Y %H:%M} UTC · Documento borrador: requiere revisión y aprobación humana.")
    for nombre, filas in secciones:
        if not filas:
            continue
        doc.add_heading(nombre, level=1)
        if len(filas[0]) > 1:
            table = doc.add_table(rows=len(filas), cols=len(filas[0]))
            table.style = "Light Grid Accent 1"
            for i, fila in enumerate(filas):
                for j, celda in enumerate(fila):
                    table.cell(i, j).text = str(celda)
        else:
            for fila in filas:
                doc.add_paragraph(fila[0])
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_xlsx(hojas: dict[str, list[list]]) -> bytes:
    """Genera un .xlsx con hojas (checklist, matriz de cumplimiento)."""
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nombre, filas in hojas.items():
        ws = wb.create_sheet(title=nombre[:31])
        for fila in filas:
            ws.append(list(fila))
        for column_cells in ws.columns:
            width = max(len(str(c.value or "")) for c in column_cells)
            ws.column_dimensions[column_cells[0].column_letter].width = min(60, max(12, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_pdf_matriz(titulo: str, filas: list[list[str]]) -> bytes:
    """Genera un PDF sencillo de la matriz de cumplimiento."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(letter), title=titulo)
    styles = getSampleStyleSheet()
    story = [Paragraph(titulo, styles["Title"]),
             Paragraph(f"Generado por LicitaIA · {datetime.utcnow():%d/%m/%Y} · Borrador sujeto a revisión humana.", styles["Normal"]),
             Spacer(1, 12)]
    data = [[Paragraph(str(c), styles["Normal"]) for c in fila] for fila in filas]
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3B67")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
    ]))
    story.append(table)
    doc.build(story)
    return buf.getvalue()
